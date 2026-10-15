#!/usr/bin/env python3
"""Pack model spectra into one spreadsheet that can be attached to a chat, and unpack it again.

On the computer that holds the spectra (needs numpy and openpyxl: pip install numpy openpyxl):

    python3 pack_model_spectra.py ~/Downloads/NLTE

writes ~/Downloads/NLTE_optical.xlsx: every text spectrum found under that folder (any depth),
cut to 3000-10000 Å.  Nothing is converted: the raw numeric columns are stored as read, with each
file's relative path, header lines, row counts and first/last raw rows, so units and columns can
be checked on the analysis side.  Of files whose first column is not a recognisable wavelength
(input profiles, light curves, ...) only the first 200 rows are kept.

On the analysis side:

    python3 -I pack_model_spectra.py --unpack NLTE_optical.xlsx --outdir data/models/NLTE

re-creates the (cut) files under --outdir with their original relative paths.

Stand-alone on purpose (no imports from this repository); the column and unit logic mirrors
bv/models.py.
"""
import argparse
import os
import re
import sys

C_AA_S = 2.99792458e18
UNIT_TO_AA = (("AA", 1.0), ("nm", 10.0), ("um", 1e4), ("cm", 1e8), ("m", 1e10))
TEXT_EXT = {".dat", ".txt", ".spec", ".spc", ".csv", ".flx", ".out", ".asc", ""}
MAX_ROWS_PER_SHEET = 1_000_000
MAX_FILE_BYTES = 50 * 1024 * 1024
UNKNOWN_KEEP = 200            # rows kept from files whose wavelength unit is not recognised


def _to_float(tok):
    return float(re.sub(r"(?<=[0-9.])[dD](?=[+-]?\d)", "e", tok))


def read_columns(path):
    """Numeric rows with the most common column count, and the non-numeric lines."""
    header, rows = [], []
    with open(path, "r", errors="replace") as fh:
        for line in fh:
            s = line.strip()
            if not s:
                continue
            if s[0] in "#!%;":
                header.append(s)
                continue
            toks = [t for t in re.split(r"[,\s;]+", s) if t]
            try:
                rows.append([_to_float(t) for t in toks])
            except ValueError:
                header.append(s)
    if not rows:
        raise ValueError("no numeric rows")
    counts = {}
    for r in rows:
        counts[len(r)] = counts.get(len(r), 0) + 1
    ncol = max(counts, key=counts.get)
    return [r for r in rows if len(r) == ncol], header, ncol, len(rows)


def infer_unit(x):
    import numpy as np
    x = np.asarray(x, float)
    x = x[np.isfinite(x) & (x > 0)]
    if x.size == 0:
        return None
    if 1e12 < np.median(x) < 1e18:
        return "Hz"
    for name, k in UNIT_TO_AA:
        lam = x * k
        if lam.min() <= 4500.0 and lam.max() >= 7000.0 and lam.max() < 1e6:
            return name
    for name, k in UNIT_TO_AA:
        if 1e3 <= np.median(x * k) <= 3e4:
            return name
    return None


def to_angstrom(x, unit):
    import numpy as np
    x = np.asarray(x, float)
    if unit == "Hz":
        with np.errstate(divide="ignore"):
            return C_AA_S / x
    return x * dict(UNIT_TO_AA)[unit]


def is_text(path):
    with open(path, "rb") as fh:
        return b"\0" not in fh.read(4096)


def pack(root, out, lmin, lmax):
    import numpy as np
    from openpyxl import Workbook

    root = os.path.abspath(os.path.expanduser(root))
    files = []
    for dirpath, dirnames, fnames in os.walk(root):
        dirnames[:] = sorted(d for d in dirnames if not d.startswith("."))
        for fn in sorted(fnames):
            p = os.path.join(dirpath, fn)
            if fn.startswith(".") or os.path.splitext(fn)[1].lower() not in TEXT_EXT:
                continue
            if os.path.getsize(p) > MAX_FILE_BYTES or not is_text(p):
                continue
            files.append(p)
    if not files:
        sys.exit(f"no text files found under {root}")

    wb = Workbook(write_only=True)
    meta = wb.create_sheet("files")
    meta.append(["file_id", "path", "n_rows", "n_kept", "n_cols", "unit", "header", "first_row", "last_row",
                 "error"])
    sheets, n_in_sheet, max_cols, kept_total = [], MAX_ROWS_PER_SHEET, 0, 0
    store = []
    for fid, p in enumerate(files, 1):
        rel = os.path.relpath(p, root)
        try:
            rows, header, ncol, nrows = read_columns(p)
        except Exception as e:
            meta.append([fid, rel, 0, 0, 0, None, None, None, None, f"{type(e).__name__}: {e}"])
            continue
        arr = np.array(rows, float)
        unit = infer_unit(arr[:, 0])
        if unit is not None:
            lam = to_angstrom(arr[:, 0], unit)
            keep = arr[(lam >= lmin) & (lam <= lmax)]
        else:                       # not recognisably a spectrum: keep a sample to look at
            keep = arr[:UNKNOWN_KEEP]
        max_cols = max(max_cols, ncol)
        meta.append([fid, rel, nrows, len(keep), ncol, unit, " | ".join(header)[:2000],
                     " ".join(repr(v) for v in rows[0]), " ".join(repr(v) for v in rows[-1]), None])
        store.append((fid, keep))
        kept_total += len(keep)
    for fid, keep in store:
        for r in keep:
            if n_in_sheet >= MAX_ROWS_PER_SHEET:
                ws = wb.create_sheet(f"data_{len(sheets) + 1}")
                ws.append(["file_id"] + [f"c{i}" for i in range(max_cols)])
                sheets.append(ws)
                n_in_sheet = 0
            sheets[-1].append([fid] + [float(v) for v in r])
            n_in_sheet += 1
    wb.save(out)
    print(f"packed {len(store)} of {len(files)} files ({kept_total} rows) from {root}")
    print(f"wrote {out} ({os.path.getsize(out) / 1e6:.1f} MB)")


def _safe_join(base, rel):
    rel = rel.replace("\\", "/")
    parts = [p for p in rel.split("/") if p not in ("", ".")]
    if not parts or any(p == ".." for p in parts) or os.path.isabs(rel):
        raise ValueError(f"unsafe path in archive: {rel!r}")
    return os.path.join(base, *parts)


def unpack(xlsx, outdir):
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    meta = {}
    rows = iter(wb["files"].iter_rows(values_only=True))
    head = next(rows)
    for r in rows:
        d = dict(zip(head, r))
        if d["file_id"] is not None and not d.get("error"):
            meta[int(d["file_id"])] = d
    data = {}
    for ws in wb.worksheets:
        if not ws.title.startswith("data_"):
            continue
        it = iter(ws.iter_rows(values_only=True))
        next(it)
        for r in it:
            fid = int(r[0])
            ncol = int(meta[fid]["n_cols"]) if fid in meta else len(r) - 1
            data.setdefault(fid, []).append(r[1:1 + ncol])
    os.makedirs(outdir, exist_ok=True)
    n = 0
    for fid, d in meta.items():
        path = _safe_join(os.path.abspath(outdir), str(d["path"]))
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w") as fh:
            for h in (d.get("header") or "").split(" | "):
                if h:
                    fh.write(h if h[0] in "#!%;" else "# " + h)
                    fh.write("\n")
            for r in data.get(fid, []):
                fh.write(" ".join(repr(float(v)) for v in r) + "\n")
        n += 1
    print(f"unpacked {n} files into {outdir}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder", nargs="?", help="folder with the model spectra (packing)")
    ap.add_argument("--out", help="output .xlsx (default: <folder>_optical.xlsx next to the folder)")
    ap.add_argument("--lmin", type=float, default=3000.0, help="shortest wavelength kept, Å")
    ap.add_argument("--lmax", type=float, default=10000.0, help="longest wavelength kept, Å")
    ap.add_argument("--unpack", metavar="XLSX", help="unpack this spreadsheet instead")
    ap.add_argument("--outdir", help="target folder for --unpack")
    a = ap.parse_args()
    if a.unpack:
        if not a.outdir:
            sys.exit("--unpack needs --outdir")
        unpack(a.unpack, a.outdir)
        return
    if not a.folder:
        ap.error("give the folder with the spectra")
    folder = os.path.abspath(os.path.expanduser(a.folder)).rstrip(os.sep)
    out = a.out or folder + "_optical.xlsx"
    pack(folder, out, a.lmin, a.lmax)


if __name__ == "__main__":
    main()
